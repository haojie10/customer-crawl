"""
广交会买家清单 (Canton Fair Buyers) 自动化采集工具
按页面和配置分卷抓取并导出为 Excel 表格
"""

import time
import httpx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

API_URL = "https://et.topease.net/api/gt8/search/getCantonFairPage"

HEADERS = {
    "authorization": "10f0e52f11bb36a9fa1d5751a5706aa6",
    "cookie": "topease_access_token=10f0e52f11bb36a9fa1d5751a5706aa6",
    "language": "en",
    "origin": "https://et.topease.net",
    "referer": "https://et.topease.net/gt8/company/result?where=search&tabCheck=9&wlf=CompanyIndex&wtm=1790059361296",
    "unionkey": "gt8",
    "content-type": "application/json",
    "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/153.0.0.0 Safari/537.36",
    "website": "https://et.topease.net"
}

BASE_PAYLOAD = {
    "autoTranslate": True,
    "buyFromArea": 0,
    "buyerCompanys": [],
    "companyNames": [],
    "companyNumbers": [],
    "companyWebsites": [],
    "containContact": "",
    "countryIds": [0],
    "dataSourceCountryId": 43,
    "excludeBank": False,
    "excludeLogistics": False,
    "excludeMicroAgent": False,
    "fullAddressMatchs": [],
    "hsCodes": [],
    "industryCategorys": [],
    "industryTyCodes": [],
    "industryType": "",
    "linkTypes": [],
    "metric": "",
    "pageSize": 100,
    "paixunt": False,
    "productDescriptions": [],
    "productTypes": [],
    "searchId": "",
    "showContactIcon": False,
    "sort": "score",
    "sortChange": False,
    "sortType": "desc",
    "supplierCompanys": [],
    "surnames": [],
    "tradeCount": -1,
    "tradeDate": 3,
    "tradeShowNameAndSessions": [],
    "tradeShowNames": [],
    "type": 9
}

COLUMNS = [
    ("#", 8),
    ("Company", 35),
    ("Nation", 18),
    ("Exhibition Session", 20),
    ("Industry", 28),
    ("Product", 25),
    ("contacts", 22),
    ("Tel", 20),
    ("Fax", 20),
    ("Address", 40),
    ("Zip Code", 14),
    ("Email", 30),
    ("Website", 30),
]

def clean_text(val):
    if val is None:
        return ""
    s = str(val).strip()
    return s

def fetch_page(client: httpx.Client, page_index: int, retry_count: int = 3):
    payload = BASE_PAYLOAD.copy()
    payload["pageIndex"] = page_index
    
    for attempt in range(retry_count):
        try:
            resp = client.post(API_URL, headers=HEADERS, json=payload, timeout=25.0)
            if resp.status_code == 200:
                data = resp.json()
                if data.get("code") == 200 or data.get("code") == 0 or "data" in data:
                    res_obj = data.get("data", {})
                    items = res_obj.get("result", [])
                    total = res_obj.get("total", 0)
                    return items, total
                else:
                    print(f"[Warn] Page {page_index} returned code {data.get('code')}: {data.get('msg')}")
            else:
                print(f"[Warn] Page {page_index} HTTP status {resp.status_code}, attempt {attempt+1}/{retry_count}")
        except Exception as e:
            print(f"[Error] Page {page_index} exception: {e}, attempt {attempt+1}/{retry_count}")
        time.sleep(1.0)
    return [], 0

def crawl_and_export(start_page: int = 1, end_page: int = 10, output_excel: str = "广交会买家_测试批次1000条.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Canton Fair Buyers"
    
    # 表头样式
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Microsoft YaHei", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # 写入表头
    headers = [col[0] for col in COLUMNS]
    ws.append(headers)
    
    for col_idx, (_, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width
    
    ws.row_dimensions[1].height = 28
    
    total_count = 0
    row_index = 1
    
    print(f"开始抓取广交会买家数据: 第 {start_page} 页 ~ 第 {end_page} 页 (每页 100 条)...")
    
    with httpx.Client() as client:
        for p in range(start_page, end_page + 1):
            print(f"--> 正在获取第 {p} / {end_page} 页...", end="", flush=True)
            records, total_records = fetch_page(client, p)
            if not records:
                print(" [空或失败]")
                continue
            
            print(f" [成功抓取 {len(records)} 条]")
            
            for item in records:
                row_index += 1
                total_count += 1
                
                # 序号
                row_data = [
                    total_count,
                    clean_text(item.get("companyName")),
                    clean_text(item.get("companyCountry")),
                    clean_text(item.get("numberSessions")),
                    clean_text(item.get("companyIndustry")),
                    clean_text(item.get("purchaseProduct")),
                    clean_text(item.get("contacts")),
                    clean_text(item.get("tel")),
                    clean_text(item.get("fax")),
                    clean_text(item.get("address")),
                    clean_text(item.get("zip")),
                    clean_text(item.get("email")),
                    clean_text(item.get("website")),
                ]
                
                ws.append(row_data)
                
                # 设置样式
                for col_idx in range(1, len(COLUMNS) + 1):
                    cell = ws.cell(row=row_index, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    if col_idx in [1, 3, 4, 11]:  # 序号、国家、展会届数、邮编居中
                        cell.alignment = center_align
                    else:
                        cell.alignment = left_align
            
            # 柔性微延时，防风控
            time.sleep(0.3)
    
    # 冻结首行
    ws.freeze_panes = "A2"
    
    wb.save(output_excel)
    print(f"\n[完成] 成功抓取并写入 {total_count} 条数据到: {output_excel}")
    return total_count

if __name__ == "__main__":
    crawl_and_export(start_page=1, end_page=10, output_excel="广交会买家_测试批次1000条.xlsx")
