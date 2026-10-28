import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
import shutil
import os

SRC_EXCEL = "欧盟采购商_近1年高频采购_有效采购决策人_(301-350页_500家公司_优耐特账号).xlsx"
BACKUP_EXCEL = "欧盟采购商_近1年高频采购_有效采购决策人_(301-350页_500家公司_优耐特账号)_原始包含未验证.xlsx"

# 1. 备份原文件
if not os.path.exists(BACKUP_EXCEL):
    shutil.copyfile(SRC_EXCEL, BACKUP_EXCEL)
    print(f"Backed up original file to: {BACKUP_EXCEL}")

wb_src = openpyxl.load_workbook(SRC_EXCEL)
ws_src = wb_src.active

COLUMNS = [
    ("#", 8),
    ("公司名称 (Company)", 35),
    ("国家/地区 (Country)", 16),
    ("注册号 (Tax Code)", 18),
    ("企业官网 (Website)", 28),
    ("最近交易日期", 15),
    ("总交易单数", 12),
    ("近3年主营/采购商品 (Top3)", 40),
    ("近3年采购HS编码 (Top3)", 30),
    ("近3年主要采购来源国", 25),
    ("近3年核心合作供应商", 40),
    ("联系人姓名", 20),
    ("职位头衔", 30),
    ("角色类型 (Purchase/Leader)", 25),
    ("电子邮箱", 35),
    ("邮箱验证状态", 20),
    ("社媒主页 (LinkedIn等)", 35),
    ("数据更新时间", 14),
]

pattern = re.compile(r"([^,\[\]\s]+)\s*\[(Valid|Risky|Unverified)\]")

companies = []
company_dict = {}

for r in range(2, ws_src.max_row + 1):
    row_vals = [ws_src.cell(row=r, column=c).value for c in range(1, 19)]
    cname = row_vals[1]
    if cname not in company_dict:
        comp_entry = {
            "company_info": row_vals[1:11],
            "valid_contacts": []
        }
        company_dict[cname] = comp_entry
        companies.append(cname)
    
    c_name = row_vals[11]
    email_val = str(row_vals[14] or "")
    if not c_name or not email_val:
        continue
    
    matches = pattern.findall(email_val)
    keep_emails = []
    keep_statuses = []
    for em, st in matches:
        if st in ["Valid", "Risky"]:
            keep_emails.append(em)
            st_text = "Valid (有效)" if st == "Valid" else "Risky (可能有效)"
            if st_text not in keep_statuses:
                keep_statuses.append(st_text)
    
    if keep_emails:
        contact_info = {
            "name": row_vals[11],
            "title": row_vals[12],
            "role": row_vals[13],
            "emails": ", ".join(keep_emails),
            "status": ", ".join(keep_statuses),
            "links": row_vals[16] or "",
            "update_date": row_vals[17] or ""
        }
        company_dict[cname]["valid_contacts"].append(contact_info)

# 2. 创建全新的精炼 Excel
wb_new = openpyxl.Workbook()
ws_new = wb_new.active
ws_new.title = "欧盟采购与企业决策人"

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
border_side = Side(border_style="thin", color="D9D9D9")
thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
data_font = Font(name="Microsoft YaHei", size=10)

headers = [col[0] for col in COLUMNS]
ws_new.append(headers)

for col_idx, (_, col_width) in enumerate(COLUMNS, 1):
    cell = ws_new.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws_new.column_dimensions[col_letter].width = col_width

ws_new.row_dimensions[1].height = 28
ws_new.freeze_panes = "A2"

row_count = 0
for cname in companies:
    comp_entry = company_dict[cname]
    cinfo = comp_entry["company_info"]
    contacts = comp_entry["valid_contacts"]
    
    if not contacts:
        # 该公司无 Valid/Risky 联系人，保留一行公司画像，联系人字段留空
        row_count += 1
        row_data = [row_count] + cinfo + ["", "", "", "", "", "", ""]
        ws_new.append(row_data)
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws_new.cell(row=row_count + 1, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if col_idx in [1, 3, 4, 6, 7, 14, 16, 18] else left_align
    else:
        for c in contacts:
            row_count += 1
            row_data = [
                row_count,
                cinfo[0], cinfo[1], cinfo[2], cinfo[3], cinfo[4],
                cinfo[5], cinfo[6], cinfo[7], cinfo[8], cinfo[9],
                c["name"], c["title"], c["role"],
                c["emails"], c["status"], c["links"], c["update_date"]
            ]
            ws_new.append(row_data)
            for col_idx in range(1, len(COLUMNS) + 1):
                cell = ws_new.cell(row=row_count + 1, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = center_align if col_idx in [1, 3, 4, 6, 7, 14, 16, 18] else left_align

wb_new.save(SRC_EXCEL)
print(f"Successfully rebuilt {SRC_EXCEL}!")
print(f"Total rows written: {row_count}")
