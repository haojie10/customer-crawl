"""
广交会买家清单 (Canton Fair Buyers) - 10万条批量采集工具
特性：
1. 自动分批持久化存储（每 50 页自动保存一次磁盘，防丢失）
2. 异常重试与断点续传
3. 动态进度展示与实时写入 Excel
"""

import os
import sys
import time
import json
import httpx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

API_URL = "https://et.topease.net/api/gt8/search/getCantonFairPage"
OUTPUT_EXCEL = "广交会买家_第01卷_前10万条(1-1000页).xlsx"
PROGRESS_FILE = "crawl_progress_100k.json"

START_PAGE = 1
END_PAGE = 1000  # 1000 页 = 100,000 条
PAGE_SIZE = 100
SAVE_INTERVAL = 50  # 每 50 页自动保存一次 Excel

HEADERS = {
    "authorization": "10f0e52f11bb36a9fa1d5751a5706aa6",
    "cookie": "topease_access_token=10f0e52f11bb36a9fa1d5751a5706aa6",
    "language": "en",
    "origin": "https://et.topease.net",
    "referer": "https://et.topease.net/gt8/company/result?where=search&tabCheck=9&wlf=CompanyIndex&wtm=1790059361296",
    "unionkey": "gt8",
    "content-type": "application/json",
    "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/153.0.0.0 Safari/537.36",
    "website": "https://et.topease.net"
}

BASE_PAYLOAD = {
    "autoTranslate": True,
    "buyFromArea": 0,
    "buyerCompanys": [],
    "companyNames": [],
    "companyNumbers": [],
    "companyWebsites": [],
    "containContact": "",
    "countryIds": [0],
    "dataSourceCountryId": 43,
    "excludeBank": False,
    "excludeLogistics": False,
    "excludeMicroAgent": False,
    "fullAddressMatchs": [],
    "hsCodes": [],
    "industryCategorys": [],
    "industryTyCodes": [],
    "industryType": "",
    "linkTypes": [],
    "metric": "",
    "pageSize": PAGE_SIZE,
    "paixunt": False,
    "productDescriptions": [],
    "productTypes": [],
    "searchId": "",
    "showContactIcon": False,
    "sort": "score",
    "sortChange": False,
    "sortType": "desc",
    "supplierCompanys": [],
    "surnames": [],
    "tradeCount": -1,
    "tradeDate": 3,
    "tradeShowNameAndSessions": [],
    "tradeShowNames": [],
    "type": 9
}

COLUMNS = [
    ("#", 8),
    ("Company", 35),
    ("Nation", 18),
    ("Exhibition Session", 20),
    ("Industry", 28),
    ("Product", 25),
    ("contacts", 22),
    ("Tel", 20),
    ("Fax", 20),
    ("Address", 40),
    ("Zip Code", 14),
    ("Email", 30),
    ("Website", 30),
]

def clean_text(val):
    if val is None:
        return ""
    return str(val).strip()

def init_excel_workbook(filepath):
    """初始化或加载现有的 Excel 文件"""
    if os.path.exists(filepath):
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        return wb, ws
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Canton Fair Buyers"
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    headers = [col[0] for col in COLUMNS]
    ws.append(headers)
    
    for col_idx, (_, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width
    
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    wb.save(filepath)
    return wb, ws

def load_progress():
    if os.path.exists(PROGRESS_FILE):
        try:
            with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {"last_completed_page": 0, "total_records": 0}

def save_progress(page: int, total_records: int):
    with open(PROGRESS_FILE, "w", encoding="utf-8") as f:
        json.dump({"last_completed_page": page, "total_records": total_records}, f, indent=2)

def fetch_page(client: httpx.Client, page_index: int, retry_count: int = 5):
    payload = BASE_PAYLOAD.copy()
    payload["pageIndex"] = page_index
    
    for attempt in range(retry_count):
        try:
            resp = client.post(API_URL, headers=HEADERS, json=payload, timeout=25.0)
            if resp.status_code == 200:
                data = resp.json()
                if data.get("code") == 200 or data.get("code") == 0 or "data" in data:
                    res_obj = data.get("data", {})
                    items = res_obj.get("result", [])
                    return items
                else:
                    print(f"\n[警告] 第 {page_index} 页接口提示: {data.get('msg')} (代码 {data.get('code')})")
            elif resp.status_code in [401, 403]:
                print(f"\n[错误] Token 可能已失效 (HTTP {resp.status_code})，请检查登录态！")
                return None
            else:
                print(f"\n[重试] 第 {page_index} 页响应状态 {resp.status_code}，重试中 ({attempt+1}/{retry_count})...")
        except Exception as e:
            print(f"\n[重试] 第 {page_index} 页异常: {e}，重试中 ({attempt+1}/{retry_count})...")
        time.sleep(1.5)
    return []

def run_crawler():
    prog = load_progress()
    last_page = prog.get("last_completed_page", 0)
    current_start = max(START_PAGE, last_page + 1)
    
    if current_start > END_PAGE:
        print(f"[提示] 目标范围 1 ~ {END_PAGE} 页已全部完成！")
        return
    
    print("=" * 60)
    print(f" 广交会买家数据采集启动 (前 10 万条)")
    print(f" 目标范围: 第 {current_start} 页 ~ 第 {END_PAGE} 页 (总计 {END_PAGE - current_start + 1} 页)")
    print(f" 输出文件: {OUTPUT_EXCEL}")
    print("=" * 60)
    
    wb, ws = init_excel_workbook(OUTPUT_EXCEL)
    data_font = Font(name="Microsoft YaHei", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    row_index = ws.max_row
    total_records = prog.get("total_records", row_index - 1)
    
    start_time = time.time()
    
    with httpx.Client() as client:
        for p in range(current_start, END_PAGE + 1):
            records = fetch_page(client, p)
            
            if records is None:
                print("\n[中断] 因认证失败中断采集。")
                break
                
            if not records:
                print(f"[警告] 第 {p} 页未获取到数据。")
                continue
            
            for item in records:
                row_index += 1
                total_records += 1
                
                row_data = [
                    total_records,
                    clean_text(item.get("companyName")),
                    clean_text(item.get("companyCountry")),
                    clean_text(item.get("numberSessions")),
                    clean_text(item.get("companyIndustry")),
                    clean_text(item.get("purchaseProduct")),
                    clean_text(item.get("contacts")),
                    clean_text(item.get("tel")),
                    clean_text(item.get("fax")),
                    clean_text(item.get("address")),
                    clean_text(item.get("zip")),
                    clean_text(item.get("email")),
                    clean_text(item.get("website")),
                ]
                
                ws.append(row_data)
                
                for col_idx in range(1, len(COLUMNS) + 1):
                    cell = ws.cell(row=row_index, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    if col_idx in [1, 3, 4, 11]:
                        cell.alignment = center_align
                    else:
                        cell.alignment = left_align
            
            # 定期存盘与进度输出
            if p % 10 == 0 or p == END_PAGE:
                elapsed = time.time() - start_time
                pages_done = p - current_start + 1
                speed = pages_done / max(1, elapsed)  # pages/sec
                remaining_pages = END_PAGE - p
                eta_minutes = (remaining_pages / max(0.01, speed)) / 60
                percent = (p / END_PAGE) * 100
                print(f"[{percent:5.1f}%] 已完成第 {p:4d}/{END_PAGE} 页 | 累计条数: {total_records:6d} | 速度: {speed:.1f}页/秒 | 预计剩余: {eta_minutes:.1f}分钟", flush=True)
            
            if p % SAVE_INTERVAL == 0 or p == END_PAGE:
                wb.save(OUTPUT_EXCEL)
                save_progress(p, total_records)
            
            # 微延时，保护账号与服务器
            time.sleep(0.2)
    
    wb.save(OUTPUT_EXCEL)
    save_progress(p, total_records)
    total_time = (time.time() - start_time) / 60
    print("\n" + "=" * 60)
    print(f"🎉 抓取完成！")
    print(f"总计获取: {total_records} 条数据")
    print(f"总耗时: {total_time:.2f} 分钟")
    print(f"文件保存至: {os.path.abspath(OUTPUT_EXCEL)}")
    print("=" * 60)

if __name__ == "__main__":
    run_crawler()
